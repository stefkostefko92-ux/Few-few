# Builds the fill-in workbook for the engineer from docs/lista-verifica-normativa.json (written by npm run lista).
# Run from argano/: python3 scripts/lista-verifica-xlsx.py  (needs openpyxl) → docs/lista-verifica-normativa.xlsx
# The summary sheet uses formulas: open the file in Excel or LibreOffice, which compute them.
import json, math
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

DOCS = Path(__file__).resolve().parent.parent / "docs"
SRC = DOCS / "lista-verifica-normativa.json"
OUT = DOCS / "lista-verifica-normativa.xlsx"
data = json.loads(SRC.read_text(encoding="utf-8"))
prof, righe = data["profilo"], data["righe"]

F = "Arial"
font = lambda **k: Font(name=F, size=k.pop("size", 10), **k)
YELLOW = PatternFill("solid", fgColor="FFFF00")
HEAD = PatternFill("solid", fgColor="1D3271")
EXAMPLE = PatternFill("solid", fgColor="EEF1F6")
thin = Side(style="thin", color="C9CFDB")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")

wb = Workbook()
# ---------- instructions ----------
ist = wb.active
ist.title = "Istruzioni"
ist.column_dimensions["A"].width = 24
ist.column_dimensions["B"].width = 100
r = 1
ist.cell(r, 1, f"Lista di verifica normativa — profilo {prof['id']}").font = font(size=14, bold=True); r += 1
ist.cell(r, 1, prof["titolo"]).font = font(size=11); r += 2
ist.cell(r, 1, "Documenti del profilo").font = font(bold=True); r += 1
for d in prof["documenti"]:
    ist.cell(r, 1, d["sigla"]).font = font(bold=True)
    c = ist.cell(r, 2, d["ambito"]); c.font = font(); c.alignment = wrap
    r += 1
r += 1
ist.cell(r, 1, "Come compilare").font = font(bold=True); r += 1
for testo in [
    "Nel foglio «Voci», per ogni riga confrontare il «Valore nel software» con il testo vigente del documento in «Dove verificare».",
    "Compilare solo le celle gialle: Esito (conforme / diverso / non applicabile), Valore corretto e Clausola esatta se diverso, Note, Verificato da, Data.",
    "La riga 2 è un ESEMPIO del formato: non è una voce e non entra nel riepilogo.",
    "Stati: «da verificare» = fonte secondaria; «stima» e «scelta del software» = decisioni del software da approvare o correggere; «derivazione» = meccanica elementare; «confermato» = due fonti indipendenti o un caso pubblicato riprodotto, da confrontare comunque con il testo; «prassi di cantiere» = indicazione di Panev Ascensori.",
    "Ogni correzione viene riportata nel registro del software (argano/src/calc/norme.ts) e cambia insieme motore di calcolo, lista e relazione di calcolo.",
    "Non trascrivere il testo delle norme in questo file: basta il valore e il numero di clausola.",
]:
    c = ist.cell(r, 2, testo); c.font = font(); c.alignment = wrap
    ist.row_dimensions[r].height = 14 * math.ceil(len(testo) / 95) + 4
    r += 1
r += 1
ist.cell(r, 1, "Legenda").font = font(bold=True)
c = ist.cell(r, 2, "celle da compilare"); c.fill = YELLOW; c.font = font(); r += 2
ist.cell(r, 1, "Verifica complessiva").font = font(bold=True); r += 1
for etichetta in ["Ingegnere (nome e cognome)", "Iscrizione all'albo", "Data", "Firma"]:
    ist.cell(r, 1, etichetta).font = font()
    c = ist.cell(r, 2, None); c.fill = YELLOW; c.border = box
    r += 1

# ---------- items ----------
ws = wb.create_sheet("Voci")
cols = [("N.", 5), ("Gruppo", 16), ("Voce", 34), ("Valore nel software", 46), ("Dove verificare", 30), ("Fonte attuale", 30), ("Stato", 14),
        ("Verifiche interessate", 30), ("Esito", 16), ("Valore corretto", 24), ("Clausola esatta", 18), ("Note", 30), ("Verificato da", 18), ("Data", 12)]
for j, (name, width) in enumerate(cols, start=1):
    c = ws.cell(1, j, name)
    c.font = font(bold=True, color="FFFFFF"); c.fill = HEAD; c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = box
    ws.column_dimensions[get_column_letter(j)].width = width
ws.row_dimensions[1].height = 30
INPUT = range(9, 15)  # Esito … Data
example = ["ESEMPIO", "Aderenza", "Coefficiente di attrito, caricamento", "μ = 0,1", "UNI EN 81-50:2020, 5.11.2", "fonti secondarie", "da verificare",
           "aderenza al caricamento", "diverso", "μ = 0,1 (con …)", "5.11.2.x", "formato della risposta: non è una voce", "iniziali", "gg/mm/aaaa"]
for j, v in enumerate(example, start=1):
    c = ws.cell(2, j, v); c.font = font(italic=True, color="5A6480"); c.fill = EXAMPLE; c.alignment = wrap; c.border = box
first = 3
for k, rw in enumerate(righe):
    rr = first + k
    voce = rw["voce"] + (f" — {rw['nota']}" if rw["nota"] else "")
    vals = [rw["n"], rw["gruppo"], voce, rw["valore"], rw["riferimento"], rw["fonte"], rw["stato"], rw["verifiche"] or "—"]
    for j, v in enumerate(vals, start=1):
        c = ws.cell(rr, j, v); c.font = font(); c.alignment = wrap; c.border = box
    for j in INPUT:
        c = ws.cell(rr, j, None); c.fill = YELLOW; c.border = box; c.font = font(color="0000FF"); c.alignment = wrap
last = first + len(righe) - 1
widths = {j: w for j, (_, w) in enumerate(cols, start=1)}
for rr in range(2, last + 1):
    lines = max(math.ceil(len(str(ws.cell(rr, j).value or "")) / max(1.0, widths[j] * 1.05)) for j in widths)
    ws.row_dimensions[rr].height = 13 * max(1, lines) + 4
dv = DataValidation(type="list", formula1='"conforme,diverso,non applicabile"', allow_blank=True, showErrorMessage=True,
                    errorTitle="Esito", error="Scegliere: conforme, diverso, non applicabile.")
ws.add_data_validation(dv)
dv.add(f"I{first}:I{last}")
ws.freeze_panes = "D3"
ws.auto_filter.ref = f"A1:N{last}"

# ---------- summary (formulas over the items, the example row left out) ----------
rs = wb.create_sheet("Riepilogo")
rs.column_dimensions["A"].width = 34
rs.column_dimensions["B"].width = 12
rs.cell(1, 1, "Riepilogo della verifica").font = font(size=13, bold=True)
rs.cell(3, 1, "Voci per stato").font = font(bold=True)
row = 4
for stato in ["da verificare", "confermato", "scelta del software", "stima", "derivazione", "prassi di cantiere"]:
    rs.cell(row, 1, stato).font = font()
    rs.cell(row, 2, f'=COUNTIF(Voci!$G${first}:$G${last},"{stato}")').font = font()
    row += 1
rs.cell(row, 1, "Totale voci").font = font(bold=True)
rs.cell(row, 2, f"=COUNTA(Voci!$A${first}:$A${last})").font = font(bold=True)
row += 2
rs.cell(row, 1, "Esito").font = font(bold=True); row += 1
for esito in ["conforme", "diverso", "non applicabile"]:
    rs.cell(row, 1, esito).font = font()
    rs.cell(row, 2, f'=COUNTIF(Voci!$I${first}:$I${last},"{esito}")').font = font()
    row += 1
rs.cell(row, 1, "da compilare").font = font()
rs.cell(row, 2, f"=COUNTBLANK(Voci!$I${first}:$I${last})").font = font()
row += 2
c = rs.cell(row, 1, f"Conteggi sulle righe {first}–{last} del foglio «Voci» (la riga 2 è l'esempio)."); c.font = font(italic=True, color="5A6480")
# print: A4 landscape, one page wide, header row repeated
for sheet in wb.worksheets:
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = sheet.page_margins.right = 0.4
ws.print_title_rows = "1:1"
wb.save(OUT)
print("saved", OUT, "rows", len(righe), "range", first, last)
